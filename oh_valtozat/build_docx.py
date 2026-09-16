# Az OH-valtozat .docx-e. Az Orvosi Hetilap megfigyelt szerkezete szerint:
# magyar Osszefoglalo + kulcsszavak, angol Summary + keywords, mindketto vegen
# az onhivatkozo "Orv Hetil. 2026; 167(x): xxx-xxx." zarosorral.
import re, pathlib
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

HERE = pathlib.Path(__file__).parent
SRC = HERE / "kezirat_szoveg.md"
OUT = HERE / "kezirat_OrvHetil.docx"

SZERZOK = [
    ("Dézsi Csilla dr.",        "1,2,3,4"),
    ("Gulyás-Oldal Viktor dr.", "1,2"),
    ("Gálóczi Imre",            "1,2"),
    ("Lábodi László dr.",       "4"),
    ("Marusin Ildikó dr.",      "4"),
    ("Szili Károly dr.",        "1,2,3,4"),
    ("Nagy Sándor dr.",         "1,2"),
]
AFFILIACIOK = [
    ("1", "Széchenyi István Egyetem, Szülészeti és Nőgyógyászati Tanszék, Győr "
          "(tanszékvezető: Nagy Sándor dr., PhD)"),
    ("2", "Széchenyi István Egyetem, Regionális- és Gazdaságtudományi Doktori Iskola "
          "(RGDI), Győr (doktori iskola vezetője: Prof. Dr. habil. Vasa László)"),
    ("3", "Szegedi Tudományegyetem, Szent-Györgyi Albert Orvostudományi Kar, "
          "Szülészeti és Nőgyógyászati Klinika, Szeged "
          "(igazgató: Prof. Dr. habil. Várbíró Szabolcs)"),
    ("4", "S.O.S. 24 Kft. – 48. Családorvosi Rendelő, Szeged "
          "(vezető: Szili Károly dr., PhD)"),
]


def alap_dokumentum():
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"
    st.font.size = Pt(12)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    pf = st.paragraph_format
    pf.line_spacing = 2.0
    pf.space_after = Pt(0)
    for s in doc.sections:
        s.top_margin = s.bottom_margin = Cm(2.5)
        s.left_margin = s.right_margin = Cm(2.5)
    return doc


def oldalszam(doc):
    for s in doc.sections:
        p = s.header.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run()
        r.font.name = "Times New Roman"; r.font.size = Pt(12)
        for instr in ("begin", "PAGE", "end"):
            e = OxmlElement("w:fldChar") if instr != "PAGE" else OxmlElement("w:instrText")
            if instr == "PAGE":
                e.set(qn("xml:space"), "preserve"); e.text = " PAGE "
            else:
                e.set(qn("w:fldCharType"), instr)
            r._r.append(e)


def cim(doc, szoveg, meret=12, felkover=True, kozep=False, utana=0):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(utana)
    if kozep:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(szoveg); r.bold = felkover; r.font.size = Pt(meret)
    return p


def bekezdes(doc, szoveg, dolt=False, meret=12, sorkoz=2.0, igazit=None):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = sorkoz
    if igazit is not None:
        p.alignment = igazit
    for i, darab in enumerate(re.split(r"\*\*(.+?)\*\*", szoveg)):
        if not darab:
            continue
        r = p.add_run(darab); r.bold = (i % 2 == 1); r.italic = dolt
        r.font.size = Pt(meret)
    return p


def uj_oldal(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def tablazat(doc, fejlec, sorok):
    t = doc.add_table(rows=1, cols=len(fejlec))
    t.style = "Table Grid"
    for i, c in enumerate(fejlec):
        cel = t.rows[0].cells[i]; cel.text = ""
        p = cel.paragraphs[0]; p.paragraph_format.line_spacing = 1.0
        r = p.add_run(c); r.bold = True; r.font.size = Pt(9); r.font.name = "Times New Roman"
    for sor in sorok:
        cellak = t.add_row().cells
        for i, c in enumerate(sor):
            cellak[i].text = ""
            p = cellak[i].paragraphs[0]; p.paragraph_format.line_spacing = 1.0
            fk = c.startswith("**") and c.endswith("**")
            r = p.add_run(c.strip("*")); r.bold = fk
            r.font.size = Pt(9); r.font.name = "Times New Roman"
    return t


szoveg = SRC.read_text(encoding="utf-8")


def szakasz(kezdet, veg=None):
    i = szoveg.index(kezdet)
    j = szoveg.index(veg) if veg else len(szoveg)
    return szoveg[i:j]


def bekezdesek(blokk):
    return [s.strip() for s in blokk.split("\n")
            if s.strip() and not s.startswith("#") and not s.startswith("---")
            and not s.startswith("|")]


doc = alap_dokumentum()
oldalszam(doc)

# ---------------------------------------------------------------- CIMOLDAL
cim(doc, "Fiatalkori elhízás és kezelése: hazai helyzetkép és a "
         "GLP-1-receptor-agonisták helye a terápiás lépcsőben",
    meret=14, kozep=True, utana=12)
bekezdes(doc, "**Rövid cím:** Fiatalkori elhízás – a terápiás lépcső",
         igazit=WD_ALIGN_PARAGRAPH.CENTER)
doc.add_paragraph()

p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.line_spacing = 1.5
for i, (nev, idx) in enumerate(SZERZOK):
    r = p.add_run(nev); r.bold = True; r.font.size = Pt(12); r.font.name = "Times New Roman"
    f = p.add_run(idx); f.font.superscript = True; f.font.size = Pt(12)
    f.font.name = "Times New Roman"
    if i < len(SZERZOK) - 1:
        v = p.add_run(", "); v.font.size = Pt(12); v.font.name = "Times New Roman"
doc.add_paragraph()

for _jel, _aff in AFFILIACIOK:
    q = doc.add_paragraph(); q.alignment = WD_ALIGN_PARAGRAPH.CENTER
    q.paragraph_format.line_spacing = 1.2; q.paragraph_format.space_after = Pt(2)
    f = q.add_run(_jel); f.font.superscript = True; f.font.size = Pt(11)
    f.font.name = "Times New Roman"
    r = q.add_run(" " + _aff); r.font.size = Pt(11); r.font.name = "Times New Roman"
doc.add_paragraph()

bekezdes(doc, "**Levelező szerző:** Szili Károly dr. · Széchenyi István Egyetem, "
              "9026 Győr, Egyetem tér 1. · szilikaroly@gmail.com",
         meret=11, sorkoz=1.2, igazit=WD_ALIGN_PARAGRAPH.CENTER)
doc.add_paragraph()
bekezdes(doc, "Beküldés előtt ellenőrizendő: a 2. affiliáció doktori iskolájának hivatalos "
              "megnevezése, valamint az Orvosi Hetilap érvényes szerzői útmutatójának "
              "terjedelmi és formai előírásai.",
         dolt=True, meret=10, sorkoz=1.2, igazit=WD_ALIGN_PARAGRAPH.CENTER)
uj_oldal(doc)

# ---------------------------------------------------------------- OSSZEFOGLALO
cim(doc, "ÖSSZEFOGLALÓ", utana=6)
for b in bekezdesek(szakasz("## ÖSSZEFOGLALÓ", "---\n\n## BEVEZETÉS")):
    if b.startswith("**Kulcsszavak"):
        doc.add_paragraph()
    bekezdes(doc, b)
uj_oldal(doc)

# ---------------------------------------------------------------- SUMMARY
cim(doc, "SUMMARY", utana=6)
for b in bekezdesek(szakasz("## SUMMARY", "---\n\n## ANYAGI TÁMOGATÁS")):
    if b.startswith("**Keywords"):
        doc.add_paragraph()
    bekezdes(doc, b)
uj_oldal(doc)

# ---------------------------------------------------------------- SZOVEG
torzs = szakasz("## BEVEZETÉS", "---\n\n## SUMMARY")
for sor in torzs.split("\n"):
    s = sor.strip()
    if not s or s.startswith("---") or s.startswith("|"):
        continue
    if s.startswith("### "):
        cim(doc, s[4:].strip(), meret=12, utana=4)
    elif s.startswith("## "):
        cim(doc, s[3:].strip(), meret=13, utana=6)
    else:
        bekezdes(doc, s)
uj_oldal(doc)

# ---------------------------------------------------------------- NYILATKOZATOK
for fejezet in ("## ANYAGI TÁMOGATÁS ÉS ÉRDEKELTSÉGEK", "## SZERZŐI MUNKAMEGOSZTÁS"):
    veg = "## SZERZŐI MUNKAMEGOSZTÁS" if "ANYAGI" in fejezet else "---\n\n## IRODALOM"
    cim(doc, fejezet[3:], utana=4)
    for b in bekezdesek(szakasz(fejezet, veg)):
        bekezdes(doc, b, dolt=b.startswith("*"))
uj_oldal(doc)

# ---------------------------------------------------------------- IRODALOM
cim(doc, "IRODALOM", utana=6)
for b in bekezdesek(szakasz("## IRODALOM", "---\n\n## TÁBLÁZATOK")):
    p = bekezdes(doc, b, meret=11, sorkoz=1.5)
    p.paragraph_format.space_after = Pt(4)
uj_oldal(doc)

# ---------------------------------------------------------------- TABLAZATOK
tabl_blokk = szakasz("## TÁBLÁZATOK", "## ÁBRÁK ÉS ÁBRAALÁÍRÁSOK")
elso = True
for resz in re.split(r"^### ", tabl_blokk, flags=re.M)[1:]:
    sorok = resz.split("\n")
    if not elso:
        uj_oldal(doc)
    elso = False
    cim(doc, sorok[0].strip(), utana=6)
    buffer = []
    for sor in sorok[1:]:
        s2 = sor.strip()
        if s2.startswith("|"):
            cellak = [c.strip() for c in s2.strip("|").split("|")]
            if set("".join(cellak)) <= set("-: "):
                continue
            buffer.append(cellak)
            continue
        if buffer:
            tablazat(doc, buffer[0], buffer[1:])
            buffer = []
            doc.add_paragraph().paragraph_format.space_after = Pt(4)
        if s2 and not s2.startswith("---"):
            tiszta = re.sub(r"\*\*(.+?)\*\*", r"\1", s2).replace("*", "")
            bekezdes(doc, tiszta, dolt=s2.startswith("*"), meret=9, sorkoz=1.0)
    if buffer:
        tablazat(doc, buffer[0], buffer[1:])
uj_oldal(doc)

# ---------------------------------------------------------------- ABRAALAIRASOK
cim(doc, "ÁBRAALÁÍRÁSOK", utana=6)
for b in bekezdesek(szakasz("## ÁBRÁK ÉS ÁBRAALÁÍRÁSOK")):
    p = bekezdes(doc, b, sorkoz=1.5)
    p.paragraph_format.space_after = Pt(8)
bekezdes(doc, "Mind a négy ábra színes; külön fájlban, 300 dpi TIFF és 600 dpi PNG "
              "formátumban kerül benyújtásra, szürkeárnyalatos változattal együtt.",
         dolt=True, meret=10)

doc.save(OUT)
print("elkeszult:", OUT)
