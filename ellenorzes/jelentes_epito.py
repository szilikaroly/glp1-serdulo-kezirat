# A 6D jelentes es a biralatok .docx valtozatanak eloallitasa a markdown forrasokbol.
import json, pathlib, re
from docx import Document
from docx.shared import Pt, Cm
from docx.oxml.ns import qn

HERE = pathlib.Path(__file__).parent
GYOKER = HERE.parent

rep = json.loads((HERE / "6d_report.json").read_text(encoding="utf-8"))
jo = sum(1 for d in rep["tetelek"]
         if all(str(v).startswith("OK") or v == "NINCS DOI"
                for k, v in d.items() if k.startswith("D")))

L = ["# 6D bibliográfiai validáció és cit-ref kereszt-ellenőrzés", "",
     "**Kézirat:** Inkretin-alapú terápiák a gyermeknőgyógyászati praxisban",
     "**Források:** Crossref REST API és NCBI PubMed (esummary)", "", "## Összegzés", "",
     f"- Ellenőrzött tételek: **{rep['hivatkozasok_szama']}**",
     f"- Mind a hét dimenzióban megfelelt: **{jo}**",
     f"- Szövegben idézett, de jegyzékből hiányzó: **{rep['idezett_de_nincs_jegyzekben'] or 'nincs'}**",
     f"- Jegyzékben szereplő, de nem idézett: **{rep['jegyzekben_van_de_nincs_idezve'] or 'nincs'}**",
     f"- A számozás az első előfordulás sorrendjét követi: "
     f"**{'igen' if rep['sorrend_helyes'] else 'NEM'}**", "",
     "A hat dimenzió: D1 DOI feloldhatósága · D2 első szerző · D3 szerzőlista (et al. szabály) · "
     "D4 folyóirat · D5 kötet · D6 oldalszám vagy cikkszám. Kiegészítő dimenzió: D7 évszám "
     "(a nyomtatott megjelenés az irányadó, mert a Crossref issued mezője gyakran az online dátum).",
     "", "## Tételes eredmény", "",
     "| # | DOI | PMID | D1 | D2 | D3 | D4 | D5 | D6 | D7 |",
     "|---|---|---|---|---|---|---|---|---|---|"]
for d in rep["tetelek"]:
    L.append(f"| {d['n']} | {d['doi'] or '–'} | {d['pmid']} | {d['D1_doi']} | {d['D2_first']} | "
             f"{d['D3_authors']} | {d['D4_journal']} | {d['D5_volume']} | {d['D6_pages']} | "
             f"{d['D7_year']} |")
L += ["", "## A futások során feloldott jelzések", "", "| Jelzés | Feloldás |", "|---|---|",
 "| kötet „eltér” füzetszám nélküli folyóiratnál | Ellenőrzési műtermék: cikkszámos folyóiratok "
 "(pl. Front Endocrinol 2022;13:897776). Az ellenőrző mintázata javítva. |",
 "| kötet „eltér” Dove Press-kiadványnál | A Crossref „Volume 19” formában adja meg a kötetet; "
 "az ellenőrző immár a PubMed értékét és csak a számjegyeket veszi. |",
 "| oldalszám „ellenőrizni” | **Valódi javítás volt:** az e1420–32 rövidített tartomány "
 "e1420–e1432 formára javítva. |",
 "| évszám „eltér” | Ellenőrzési műtermék: a Crossref issued az online megjelenés dátuma; "
 "a nyomtatott év az irányadó. |",
 "| első szerző és „et al.” | Ellenőrzési műtermék: testületi szerző. |",
 "| nincs DOI | A Clin Exp Obstet Gynecol 2011-es közleményének nincs DOI-ja; PMID alapján ellenőrizve. |",
 "| **az egész tábla „eltér”** | A DOI/PMID-térkép elavult a hivatkozásszám változása után. "
 "Javítva: a térkép az `ids.json`-ból **származtatott**, a sorszámozó generálja. |",
 "", "## Újrafuttatás", "", "```bash", "python osszefuz_es_szamoz.py",
 "python ellenorzes/validate_6d.py", "```"]
(HERE / "6d_jelentes.md").write_text("\n".join(L), encoding="utf-8")


def md2docx(src, out, cimlap):
    t = pathlib.Path(src).read_text(encoding="utf-8")
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"
    st.font.size = Pt(11)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    st.paragraph_format.line_spacing = 1.3
    st.paragraph_format.space_after = Pt(6)
    for s in doc.sections:
        s.top_margin = s.bottom_margin = Cm(2.5)
        s.left_margin = s.right_margin = Cm(2.0)
    for raw in t.split("\n"):
        line = raw.rstrip()
        if not line or line.strip() == "---":
            continue
        if line.startswith("|"):
            cells = [c.strip() for c in line.strip("|").split("|")]
            if set("".join(cells)) <= set("-: "):
                continue
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(" | ".join(cells))
            r.font.size = Pt(8.5)
            continue
        lvl = len(line) - len(line.lstrip("#"))
        txt = line.lstrip("# ").strip()
        if lvl:
            h = doc.add_paragraph()
            h.paragraph_format.space_before = Pt(12 if lvl < 3 else 8)
            r = h.add_run(txt)
            r.bold = True
            r.font.size = Pt(14 - lvl)
            continue
        p = doc.add_paragraph()
        for i, part in enumerate(re.split(r"\*\*(.+?)\*\*", txt)):
            r = p.add_run(part)
            r.bold = (i % 2 == 1)
    doc.core_properties.title = cimlap
    doc.save(out)
    print("kesz:", out)


md2docx(HERE / "6d_jelentes.md", HERE / "6d_jelentes.docx", "6D bibliográfiai validáció")
for nev, cimlap in [("1_rigorozus_biralat", "Rigorózus bírálat"),
                    ("2_metodikai_elettani_biokemiai_biralat", "Metodikai–élettani–biokémiai bírálat"),
                    ("3_javitasi_naplo", "Javítási napló")]:
    md2docx(GYOKER / "biralatok" / f"{nev}.md", GYOKER / "biralatok" / f"{nev}.docx", cimlap)
